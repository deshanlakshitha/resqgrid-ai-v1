"""Export the two documentation sources to PDF/DOCX and verify their contents.

Run from any directory: python docs/generate_documents.py
Documentation-only dependencies: python-docx, reportlab, PyMuPDF (validation).
No application import, environment-file access, or network request is performed.
"""

from __future__ import annotations

import argparse
import re
import textwrap
from dataclasses import dataclass, field
from html import escape
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor
from reportlab.lib import colors
from reportlab.lib.enums import TA_LEFT
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import (
    HRFlowable, LongTable, PageBreak, Paragraph, Preformatted,
    SimpleDocTemplate, Spacer, TableStyle,
)
from reportlab.platypus.tableofcontents import TableOfContents

ROOT = Path(__file__).resolve().parent
EXPORTS = ROOT / "exports"
BASELINE = "0d85c00"
NAVY = "12243B"
TEAL = "087F8C"
GRAY = "526174"
PALE = "EDF5F7"
SOURCES = (
    ("PROJECT_DOCUMENTATION.md", "ResQGrid-AI-Complete-Guide", "COMPLETE PROJECT GUIDE"),
    ("DOCUMENTARY_SCRIPT.md", "ResQGrid-AI-Documentary-Script", "DOCUMENTARY PRODUCTION SCRIPT"),
)
TOKEN = re.compile(r"\[([^\]]+)\]\(([^)]+)\)|`([^`]+)`|\*\*([^*]+)\*\*")


@dataclass
class Block:
    kind: str
    text: str = ""
    level: int = 0
    rows: list[list[str]] = field(default_factory=list)


def parse_markdown(path: Path) -> list[Block]:
    lines = path.read_text(encoding="utf-8").splitlines()
    blocks: list[Block] = []
    i = 0
    while i < len(lines):
        line = lines[i]
        if not line.strip():
            i += 1
            continue
        if line.startswith("```"):
            code = []
            i += 1
            while i < len(lines) and not lines[i].startswith("```"):
                code.append(lines[i])
                i += 1
            if i == len(lines):
                raise ValueError(f"Unclosed code fence: {path.name}")
            blocks.append(Block("code", "\n".join(code)))
            i += 1
            continue
        heading = re.match(r"^(#{1,3}) (.+)$", line)
        if heading:
            blocks.append(Block("heading", heading[2], len(heading[1])))
            i += 1
            continue
        if line.startswith("|"):
            rows = []
            while i < len(lines) and lines[i].startswith("|"):
                cells = [c.strip() for c in lines[i].strip().strip("|").split("|")]
                if not all(re.fullmatch(r":?-{3,}:?", c) for c in cells):
                    rows.append(cells)
                i += 1
            if len({len(row) for row in rows}) != 1:
                raise ValueError(f"Inconsistent table columns: {path.name}")
            blocks.append(Block("table", rows=rows))
            continue
        if line.startswith("> "):
            blocks.append(Block("quote", line[2:]))
            i += 1
            continue
        if line.startswith("- "):
            blocks.append(Block("bullet", line[2:]))
            i += 1
            continue
        paragraph = [line]
        i += 1
        while i < len(lines) and lines[i].strip() and not re.match(r"^(#|>|\||- |```)", lines[i]):
            paragraph.append(lines[i])
            i += 1
        blocks.append(Block("paragraph", " ".join(paragraph)))
    return blocks


def plain(text: str) -> str:
    return TOKEN.sub(lambda m: m[1] or m[3] or m[4], text)


def inline_pdf(text: str, mono_size: float = 8.2) -> str:
    parts = []
    position = 0
    for match in TOKEN.finditer(text):
        parts.append(escape(text[position:match.start()]))
        if match[1]:
            parts.append(f'<link href="{escape(match[2], quote=True)}" color="#{TEAL}">{escape(match[1])}</link>')
        elif match[3]:
            parts.append(f'<font name="DocMono" size="{mono_size}">{escape(match[3])}</font>')
        else:
            parts.append(f"<b>{escape(match[4])}</b>")
        position = match.end()
    parts.append(escape(text[position:]))
    return "".join(parts)


def pdf_color(value: str):
    return colors.HexColor("#" + value.lstrip("#"))


def register_fonts() -> None:
    candidates = (
        (Path("C:/Windows/Fonts"), ("arial.ttf", "arialbd.ttf", "ariali.ttf", "arialbi.ttf", "consola.ttf")),
        (Path("/usr/share/fonts/truetype/dejavu"), ("DejaVuSans.ttf", "DejaVuSans-Bold.ttf", "DejaVuSans-Oblique.ttf", "DejaVuSans-BoldOblique.ttf", "DejaVuSansMono.ttf")),
    )
    for directory, names in candidates:
        if all((directory / name).is_file() for name in names):
            for family, name in zip(("DocSans", "DocSansBold", "DocSansItalic", "DocSansBoldItalic", "DocMono"), names):
                pdfmetrics.registerFont(TTFont(family, str(directory / name)))
            pdfmetrics.registerFontFamily("DocSans", normal="DocSans", bold="DocSansBold", italic="DocSansItalic", boldItalic="DocSansBoldItalic")
            return
    raise RuntimeError("Install Arial/Consolas on Windows or DejaVu Sans/Mono on Linux for Unicode PDF export.")


def wrapped_code(text: str) -> str:
    result = []
    for line in text.splitlines():
        indent = len(line) - len(line.lstrip())
        result.extend(textwrap.wrap(line, width=100, subsequent_indent=" " * min(indent + 4, 16),
                                    replace_whitespace=False, drop_whitespace=False,
                                    break_long_words=True, break_on_hyphens=False) or [""])
    return "\n".join(result)


def column_weights(rows: list[list[str]]) -> list[float]:
    header = rows[0]
    if header == ["Method", "Path", "Access", "Purpose"]:
        return [0.10, 0.34, 0.16, 0.40]
    if header[0] == "Operation" and len(header) == 4:
        return [0.27, 0.15, 0.29, 0.29]
    if header == ["Role", "Email", "Demo password"]:
        return [0.23, 0.50, 0.27]
    if len(header) == 2:
        return [0.30, 0.70]
    if header[0] == "Setting":
        return [0.33, 0.40, 0.27]
    if header[0] == "Scene":
        return [0.09, 0.18, 0.29, 0.44]
    lengths = [sum(min(len(plain(row[c])), 90) for row in rows) / len(rows) for c in range(len(header))]
    weights = [max(length, 10) ** 0.65 for length in lengths]
    return [weight / sum(weights) for weight in weights]


class GuidePDF(SimpleDocTemplate):
    def afterFlowable(self, flowable):
        if isinstance(flowable, Paragraph) and hasattr(flowable, "toc_key"):
            key = flowable.toc_key
            self.canv.bookmarkPage(key)
            self.canv.addOutlineEntry(flowable.getPlainText(), key, level=0, closed=False)
            self.notify("TOCEntry", (0, flowable.getPlainText(), self.page, key))


def pdf_export(blocks: list[Block], destination: Path, label: str) -> None:
    width, height = A4
    margin = 52
    usable = width - margin * 2
    base = dict(fontName="DocSans", textColor=pdf_color(NAVY), alignment=TA_LEFT)
    styles = {
        "body": ParagraphStyle("body", fontSize=10, leading=14.4, spaceAfter=8, splitLongWords=True, **base),
        "quote": ParagraphStyle("quote", fontSize=10, leading=14.4, backColor=pdf_color(PALE),
                                borderPadding=10, leftIndent=10, rightIndent=10, spaceBefore=8, spaceAfter=12, **base),
        "h2": ParagraphStyle("h2", fontSize=20, leading=25, fontName="DocSansBold", textColor=pdf_color(NAVY), spaceAfter=16, keepWithNext=True),
        "h3": ParagraphStyle("h3", fontSize=12, leading=16, fontName="DocSansBold", textColor=pdf_color(TEAL), spaceBefore=10, spaceAfter=7, keepWithNext=True),
        "cell": ParagraphStyle("cell", fontSize=8.2, leading=11.4, splitLongWords=True, **base),
        "header": ParagraphStyle("header", fontSize=8.3, leading=11.5, fontName="DocSansBold", textColor=colors.white),
        "code": ParagraphStyle("code", fontName="DocMono", fontSize=7.6, leading=10.5,
                               textColor=pdf_color(NAVY), backColor=pdf_color("F1F4F8"),
                               borderPadding=8, spaceBefore=5, spaceAfter=12),
        "bullet": ParagraphStyle("bullet", fontSize=10, leading=14.4, leftIndent=12, firstLineIndent=-10, spaceAfter=6, **base),
    }
    story = []
    first_chapter = next(i for i, b in enumerate(blocks) if b.kind == "heading" and b.level == 2)
    title_style = ParagraphStyle("title", fontName="DocSansBold", fontSize=32, leading=39, textColor=pdf_color(NAVY), spaceAfter=23)
    label_style = ParagraphStyle("label", fontName="DocSansBold", fontSize=10, leading=14, textColor=pdf_color(TEAL), spaceAfter=26)
    story.extend([Spacer(1, 65), Paragraph(label, label_style), Paragraph(inline_pdf(blocks[0].text), title_style),
                  HRFlowable(width="100%", thickness=3, color=pdf_color(TEAL)), Spacer(1, 22)])
    for block in blocks[1:first_chapter]:
        story.append(Paragraph(inline_pdf(block.text), styles["quote" if block.kind == "quote" else "body"]))
    story.extend([Spacer(1, 28), Paragraph("SOURCE-BASED EDITION · FICTIONAL DEMONSTRATION DATA ONLY", label_style), PageBreak()])
    story.append(Paragraph("Contents", styles["h2"]))
    toc = TableOfContents()
    toc.levelStyles = [ParagraphStyle("toc", fontName="DocSans", fontSize=9.4, leading=15.3,
                                     textColor=pdf_color(NAVY), leftIndent=0, firstLineIndent=0, spaceBefore=2)]
    toc.dotsMinLevel = 0
    story.append(toc)
    chapter = 0
    for block in blocks[first_chapter:]:
        if block.kind == "heading":
            if block.level == 2:
                story.append(PageBreak())
                chapter += 1
                paragraph = Paragraph(inline_pdf(block.text), styles["h2"])
                paragraph.toc_key = f"chapter-{chapter}"
                story.append(paragraph)
            else:
                story.append(Paragraph(inline_pdf(block.text), styles["h3"]))
        elif block.kind == "table":
            data = []
            for row_index, row in enumerate(block.rows):
                style = styles["header" if row_index == 0 else "cell"]
                data.append([Paragraph(inline_pdf(cell, 7.5), style) for cell in row])
            table = LongTable(data, colWidths=[w * usable for w in column_weights(block.rows)], repeatRows=1, hAlign="LEFT")
            table.setStyle(TableStyle([
                ("BACKGROUND", (0, 0), (-1, 0), pdf_color(NAVY)),
                ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, pdf_color("F2F6FA")]),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("LEFTPADDING", (0, 0), (-1, -1), 7), ("RIGHTPADDING", (0, 0), (-1, -1), 7),
                ("TOPPADDING", (0, 0), (-1, -1), 7), ("BOTTOMPADDING", (0, 0), (-1, -1), 7),
                ("LINEBELOW", (0, 0), (-1, 0), 0.7, pdf_color(TEAL)),
                ("LINEBELOW", (0, 1), (-1, -1), 0.25, pdf_color("DCE4EC")),
            ]))
            story.extend([table, Spacer(1, 12)])
        elif block.kind == "code":
            story.append(Preformatted(wrapped_code(block.text), styles["code"]))
        else:
            kind = block.kind if block.kind in {"quote", "bullet"} else "body"
            content = ("• " if block.kind == "bullet" else "") + inline_pdf(block.text)
            story.append(Paragraph(content, styles[kind]))

    def page_decoration(canvas, doc):
        canvas.saveState()
        if doc.page > 1:
            canvas.setFont("DocSans", 8)
            canvas.setFillColor(pdf_color(GRAY))
            canvas.drawString(margin, height - 29, "RESQGRID AI")
            canvas.drawRightString(width - margin, height - 29, label)
            canvas.setStrokeColor(pdf_color("DCE4EC"))
            canvas.line(margin, height - 36, width - margin, height - 36)
        canvas.setFont("DocSans", 7.5)
        canvas.setFillColor(pdf_color(GRAY))
        canvas.drawString(margin, 29, f"Source baseline {BASELINE}  |  Demonstration only")
        canvas.drawRightString(width - margin, 29, f"Page {doc.page}")
        canvas.restoreState()

    doc = GuidePDF(str(destination), pagesize=A4, leftMargin=margin, rightMargin=margin,
                   topMargin=55, bottomMargin=49, title=blocks[0].text, author="ResQGrid AI Team",
                   subject=label, pageCompression=1)
    doc.multiBuild(story, onFirstPage=page_decoration, onLaterPages=page_decoration)


def hyperlink(paragraph, label: str, target: str, internal: bool = False) -> None:
    link = OxmlElement("w:hyperlink")
    if internal:
        link.set(qn("w:anchor"), target)
    else:
        from docx.opc.constants import RELATIONSHIP_TYPE
        relationship = paragraph.part.relate_to(target, RELATIONSHIP_TYPE.HYPERLINK, is_external=True)
        link.set(qn("r:id"), relationship)
    run = OxmlElement("w:r")
    props = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), TEAL)
    props.append(color)
    run.append(props)
    text = OxmlElement("w:t")
    text.text = label
    run.append(text)
    link.append(run)
    paragraph._p.append(link)


def inline_docx(paragraph, text: str) -> None:
    position = 0
    for match in TOKEN.finditer(text):
        paragraph.add_run(text[position:match.start()])
        if match[1]:
            hyperlink(paragraph, match[1], match[2])
        elif match[3]:
            run = paragraph.add_run(match[3])
            run.font.name = "Consolas"
            run.font.size = Pt(8.5)
        else:
            paragraph.add_run(match[4]).bold = True
        position = match.end()
    paragraph.add_run(text[position:])


def shade(element, color: str) -> None:
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), color)
    element.append(shading)


def docx_export(blocks: list[Block], destination: Path, label: str) -> None:
    document = Document()
    document.core_properties.title = blocks[0].text
    document.core_properties.author = "ResQGrid AI Team"
    document.core_properties.subject = label
    document.core_properties.version = "1.0"
    section = document.sections[0]
    section.page_width, section.page_height = Mm(210), Mm(297)
    section.top_margin, section.bottom_margin = Mm(20), Mm(18)
    section.left_margin = section.right_margin = Mm(19)
    section.header_distance = section.footer_distance = Mm(9)
    for name in ("Normal", "Body Text", "List Bullet"):
        style = document.styles[name]
        style.font.name = "Arial"
        style.font.size = Pt(10)
        style.font.color.rgb = RGBColor.from_string(NAVY)
        style.paragraph_format.line_spacing = 1.16
        style.paragraph_format.space_after = Pt(7)
    for name, size, color in (("Title", 32, NAVY), ("Heading 1", 20, NAVY), ("Heading 2", 12, TEAL)):
        style = document.styles[name]
        style.font.name = "Arial"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_after = Pt(12)
        style.paragraph_format.keep_with_next = True
    section.header.paragraphs[0].text = "RESQGRID AI  |  " + label
    section.header.paragraphs[0].runs[0].font.size = Pt(8)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer.add_run(f"Source {BASELINE}  |  Demonstration only  |  Page ").font.size = Pt(8)
    page_field = OxmlElement("w:fldSimple")
    page_field.set(qn("w:instr"), "PAGE")
    footer._p.append(page_field)

    first_chapter = next(i for i, b in enumerate(blocks) if b.kind == "heading" and b.level == 2)
    document.add_paragraph().paragraph_format.space_after = Pt(55)
    p = document.add_paragraph(label)
    p.runs[0].bold = True
    p.runs[0].font.color.rgb = RGBColor.from_string(TEAL)
    document.add_paragraph(blocks[0].text, "Title")
    for block in blocks[1:first_chapter]:
        p = document.add_paragraph()
        inline_docx(p, block.text)
        if block.kind == "quote":
            shade(p._p.get_or_add_pPr(), PALE)
            p.paragraph_format.space_before = Pt(15)
            p.paragraph_format.left_indent = Mm(3)
    document.add_paragraph("SOURCE-BASED EDITION · FICTIONAL DEMONSTRATION DATA ONLY")
    document.add_page_break()
    document.add_paragraph("Contents", "Heading 1")
    chapters = [b for b in blocks if b.kind == "heading" and b.level == 2]
    for index, chapter in enumerate(chapters, 1):
        p = document.add_paragraph()
        p.paragraph_format.space_after = Pt(3)
        hyperlink(p, chapter.text, f"chapter_{index}", internal=True)

    chapter_index = 0
    for block in blocks[first_chapter:]:
        if block.kind == "heading":
            if block.level == 2:
                chapter_index += 1
                p = document.add_paragraph(block.text, "Heading 1")
                p.paragraph_format.page_break_before = True
                start = OxmlElement("w:bookmarkStart")
                start.set(qn("w:id"), str(chapter_index))
                start.set(qn("w:name"), f"chapter_{chapter_index}")
                end = OxmlElement("w:bookmarkEnd")
                end.set(qn("w:id"), str(chapter_index))
                p._p.insert(0, start)
                p._p.append(end)
            else:
                document.add_paragraph(block.text, "Heading 2")
        elif block.kind == "table":
            table = document.add_table(rows=0, cols=len(block.rows[0]))
            table.autofit = False
            weights = column_weights(block.rows)
            for column, weight in zip(table.columns, weights):
                column.width = Mm(172 * weight)
            for row_index, row in enumerate(block.rows):
                cells = table.add_row().cells
                row_props = table.rows[-1]._tr.get_or_add_trPr()
                row_props.append(OxmlElement("w:cantSplit"))
                if row_index == 0:
                    row_props.append(OxmlElement("w:tblHeader"))
                for cell, text, weight in zip(cells, row, weights):
                    cell.width = Mm(172 * weight)
                    shade(cell._tc.get_or_add_tcPr(), NAVY if row_index == 0 else ("F2F6FA" if row_index % 2 == 0 else "FFFFFF"))
                    p = cell.paragraphs[0]
                    p.paragraph_format.space_before = Pt(4)
                    p.paragraph_format.space_after = Pt(5)
                    p.paragraph_format.line_spacing = 1.1
                    inline_docx(p, text)
                    for run in p.runs:
                        run.font.size = Pt(8.3)
                        if row_index == 0:
                            run.bold = True
                            run.font.color.rgb = RGBColor(255, 255, 255)
            document.add_paragraph().paragraph_format.space_after = Pt(2)
        elif block.kind == "code":
            p = document.add_paragraph()
            shade(p._p.get_or_add_pPr(), "F1F4F8")
            p.paragraph_format.left_indent = Mm(2)
            p.paragraph_format.line_spacing = 1.05
            p.paragraph_format.keep_together = True
            run = p.add_run(wrapped_code(block.text))
            run.font.name = "Consolas"
            run.font.size = Pt(8)
        else:
            p = document.add_paragraph(style="List Bullet" if block.kind == "bullet" else "Normal")
            inline_docx(p, block.text)
            if block.kind == "quote":
                shade(p._p.get_or_add_pPr(), PALE)
                p.paragraph_format.left_indent = Mm(3)
    document.save(destination)


def normalized(text: str) -> str:
    return re.sub(r"\s+", "", text).replace("\u00ad", "")


def validate(blocks: list[Block], basename: str) -> dict:
    try:
        import fitz
    except ImportError as exc:
        raise RuntimeError("Install PyMuPDF for export validation: python -m pip install PyMuPDF") from exc
    pdf_path = EXPORTS / f"{basename}.pdf"
    word_path = EXPORTS / f"{basename}.docx"
    for path in (pdf_path, word_path):
        if not path.is_file() or path.stat().st_size < 1000:
            raise AssertionError(f"Missing/empty export: {path.name}")
    word = Document(word_path)
    word_text = normalized(" ".join(word.element.xpath("//w:t/text()")))
    with fitz.open(pdf_path) as pdf:
        pdf_text = normalized("\n".join(
            page.get_text(sort=False, clip=fitz.Rect(45, 45, page.rect.width - 45, page.rect.height - 45))
            for page in pdf
        ))
        expected = []
        for block in blocks:
            if block.kind == "table":
                expected.extend(plain(cell) for row in block.rows for cell in row)
            else:
                expected.append(block.text if block.kind == "code" else plain(block.text))
        for label, content in (("DOCX", word_text), ("PDF", pdf_text)):
            missing = [text[:100] for text in expected if normalized(text) not in content]
            if missing:
                raise AssertionError(f"{basename} {label}: missing source content: {missing[:5]}")
        overflow = []
        for index, page in enumerate(pdf, 1):
            for block in page.get_text("dict")["blocks"]:
                for line in block.get("lines", []):
                    for span in line["spans"]:
                        x0, y0, x1, y1 = span["bbox"]
                        if x0 < 45 or x1 > page.rect.width - 45 or y0 < 12 or y1 > page.rect.height - 12:
                            overflow.append((index, span["text"][:65]))
        if overflow:
            raise AssertionError(f"Text outside safe page bounds: {overflow[:8]}")
        chapters = sum(b.kind == "heading" and b.level == 2 for b in blocks)
        if len(pdf.get_toc()) != chapters:
            raise AssertionError("PDF chapter bookmarks are incomplete")
        result = {"document": basename, "pages": len(pdf), "chapters": chapters,
                  "source_blocks_checked": len(expected), "tables": len(word.tables),
                  "pdf_bytes": pdf_path.stat().st_size, "docx_bytes": word_path.stat().st_size}
    return result


def preview(basename: str, destination: Path) -> None:
    import fitz
    from PIL import Image, ImageDraw
    destination.mkdir(parents=True, exist_ok=True)
    with fitz.open(EXPORTS / f"{basename}.pdf") as pdf:
        indices = sorted({0, 1, 2, len(pdf) // 3, len(pdf) // 2, len(pdf) - 1})
        thumbnails = []
        for index in indices:
            pix = pdf[index].get_pixmap(matrix=fitz.Matrix(0.72, 0.72), alpha=False)
            image = Image.frombytes("RGB", (pix.width, pix.height), pix.samples)
            thumbnails.append((index, image))
        tile_width = max(image.width for _, image in thumbnails) + 20
        tile_height = max(image.height for _, image in thumbnails) + 35
        sheet = Image.new("RGB", (tile_width * 3, tile_height * 2), "#dce4ec")
        draw = ImageDraw.Draw(sheet)
        for n, (index, image) in enumerate(thumbnails):
            x, y = (n % 3) * tile_width + 10, (n // 3) * tile_height + 25
            sheet.paste(image, (x, y))
            draw.text((x, y - 18), f"Page {index + 1}", fill="#12243b")
        sheet.save(destination / f"{basename}-preview.png")


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--check", action="store_true", help="Validate existing exports without regenerating")
    parser.add_argument("--preview-dir", type=Path, help="Optional local PDF contact-sheet output directory")
    args = parser.parse_args()
    if not args.check:
        EXPORTS.mkdir(exist_ok=True)
        register_fonts()
    for source, basename, label in SOURCES:
        blocks = parse_markdown(ROOT / source)
        if not args.check:
            pdf_export(blocks, EXPORTS / f"{basename}.pdf", label)
            docx_export(blocks, EXPORTS / f"{basename}.docx", label)
        print(validate(blocks, basename))
        if args.preview_dir:
            preview(basename, args.preview_dir.resolve())


if __name__ == "__main__":
    main()
