# -*- coding: utf-8 -*-
"""Generate the ResQGrid AI Pitch Day Playbook as .docx and .pdf."""

import os

OUT_DIR = os.path.dirname(os.path.abspath(__file__))
DOCX_PATH = os.path.join(OUT_DIR, "ResQGrid-AI-Pitch-Playbook.docx")
PDF_PATH = os.path.join(OUT_DIR, "ResQGrid-AI-Pitch-Playbook.pdf")

# ---------------------------------------------------------------- content ---
# block kinds: title, subtitle, meta, h1, h2, h3, p, say, dir, b, table, box, pb

CONTENT = [
    ("title", "ResQGrid AI — Pitch Day Playbook"),
    ("subtitle", "Intelligent Emergency Resource Network"),
    ("meta", [
        "Event:  AI Buildathon — Pitch Day, Wednesday, 16 September 2026",
        "Venue:  Tilapiya Colombo, Race Course Avenue, Colombo 07",
        "Format:  10 minutes pitch + live demo (warning at 8:00, hard stop at 10:00) + 3 minutes Q&A",
        "Product line:  “AI recommends. Humans approve. Every important decision is explainable and auditable.”",
        "Live app:  https://web-two-mu-edisuy5tdl.vercel.app",
        "Backend:  https://resqgrid-api-9wns.onrender.com/api/v1",
        "Code:  https://github.com/deshanlakshitha/resqgrid-ai",
    ]),
    ("p", "How to use this document: Sections 1–2 map the judges’ scoring rubric to your exact 10-minute script. Section 3 is the click-by-click live demo run-sheet. Section 4 is your explainable AI-usage answer. Sections 5–6 arm you for impact claims and Q&A. Sections 7–8 are the checklists to print."),

    ("h1", "1. How You Will Be Scored — and How ResQGrid Answers"),
    ("p", "Judges award 100 points across six areas. Technical substance — AI centrality (25) plus working prototype (25) — is half of the total score, more than the story and the slides combined. The plan below spends time accordingly: the demo gets 5 of your 10 minutes."),
    ("table", {
        "headers": ["Scoring area", "Pts", "What judges look for", "How ResQGrid answers"],
        "rows": [
            ["Problem & originality", "15", "A real, named user — not a generic “AI for X”", "Named user: the duty officer at a Sri Lanka District Emergency Operations Centre (DMC). Phone calls, WhatsApp voice notes and radio chatter arrive unstructured during floods; nobody prioritizes. Most teams pitch chatbots — you pitch an operational decision system."],
            ["AI centrality & technical depth", "25", "AI does the core work, not decoration", "Qwen (Alibaba Model Studio) triage turns free text into a typed incident schema; computer-vision evidence signals; a command assistant answers from the live database; AI is load-bearing — remove it and the product stops working."],
            ["Code quality & working prototype", "25", "A real build that runs today", "Live deployment: Next.js + FastAPI + PostgreSQL/PostGIS + Redis; JWT with 4-role RBAC; Alembic migrations; Docker Compose; tests; installable PWA and a real Android APK (Capacitor). Judges can log in themselves."],
            ["Explainable AI usage", "10", "Explain how AI was used to build it, and how the product’s AI decides", "Both layers ready: (a) build story — Qoder-generated triage adapter with schema validation, reviewed and tested; (b) runtime explainability — every priority score decomposes into weighted factors with reason codes."],
            ["Real-world impact & feasibility", "15", "Numbers with sources; a credible path to real use", "330+ deaths in the Nov 2025 Sri Lanka floods (BBC); 300,000+ affected in the 2016 floods (IOM/UN sitreps); World Bank: flood damage every year for 40 years. Already deployed on free-tier cloud — near-zero MVP cost; Sinhala/Tamil and a DMC pilot on the roadmap."],
            ["Demo & presentation clarity", "10", "The real product doing real work on screen, on time", "A 5-minute scripted demo narrated at every step; warmed backend; backup video cued; rehearsed to fit the 10-minute hard stop."],
        ],
    }),

    ("h1", "2. Your 10 Minutes, Second by Second"),

    ("h2", "Segment 1 — The Problem (0:00–1:30)"),
    ("dir", "Keep one slide up: the six questions a duty officer must answer. No feature screenshots yet."),
    ("say", "In November last year, floods and mudslides in Sri Lanka killed more than 330 people. In the 2016 monsoon floods, over 300,000 people were affected — many of them around this city. And the hardest part for the people in charge was not the water. It was the chaos of information."),
    ("say", "Meet our named user: the duty officer at a District Emergency Operations Centre — the DMC officer who runs a district’s disaster response from one room. During a flood, reports reach him as phone calls, WhatsApp messages and radio chatter: free text, no structure, no priority. Within minutes he must answer six questions: Which incidents are most urgent? Who is at risk? What resources are available? Which resource goes where? Which routes are still open? And what changed since my last decision?"),
    ("say", "Today he answers with paper logs, phone calls and gut feel — while the situation changes every twenty minutes. Existing tools are static registries or chat groups. They record information. They do not understand it, rank it, or act on it. That is the problem ResQGrid AI solves."),

    ("h2", "Segment 2 — The Solution in One Breath (1:30–2:00)"),
    ("say", "ResQGrid AI is an intelligent emergency resource network. Citizens and operators report incidents in plain language. AI structures and triages them, a transparent priority engine ranks them, resource matching proposes the right unit with the safest route — and a human approves every dispatch. Our principle: AI recommends. Humans approve. Every important decision is explainable and auditable. Let me show it running — live, on the real deployment."),

    ("h2", "Segment 3 — Live Demo (2:00–7:00)"),
    ("say", "This is the live system — frontend on Vercel, API on Render, real PostgreSQL/PostGIS database. Everything you see, you can touch yourselves afterwards."),
    ("p", "Follow the click-by-click run-sheet in Section 3. Narrate every step; never demo in silence."),

    ("h2", "Segment 4 — How We Built It (7:00–8:30)"),
    ("say", "How is this built? A Next.js 14 command dashboard and a FastAPI backend, on PostgreSQL with PostGIS for geospatial search and Redis for caching. AI runs on Alibaba Cloud Model Studio — Qwen — behind a provider adapter, with a Gemini adapter as backup; if no AI key is configured at all, the system degrades gracefully to rule-based triage instead of dying. That adapter pattern is one concrete part I can walk through in detail."),
    ("say", "Three engineering decisions define the system. First: the AI never writes to the database directly — every output is validated against a strict Pydantic schema, and when information is missing the system records “unknown” instead of inventing data. Second: prioritization is deterministic and explainable — a weighted score across life risk, medical urgency, people at risk, environment, time sensitivity and evidence confidence — every point decomposable, no black box. Third: human-in-the-loop is enforced in code — a recommendation is only a proposal until an authorized operator approves it, and every AI output and human decision lands in an audit log."),

    ("h2", "Segment 5 — Impact, Feasibility, Next Steps (8:30–9:30)"),
    ("say", "Why does this matter? Sri Lanka faces flood damage every single year — the World Bank confirms four decades of it. In November 2025, more than 330 people died; in 2016, over 300,000 were affected. The lever we can pull is the speed of triage: ResQGrid structures and prioritizes a report in seconds, at a scale no duty officer can match by phone."),
    ("say", "And it is feasible now — not in a roadmap slide. The system you just saw is deployed today on free-tier cloud, so an MVP costs almost nothing to run. It is open-source, starts with one Docker Compose command, and field responders need only an Android phone — the same app installs as a PWA or a native APK. Next steps: a pilot with a District DMC office, Sinhala and Tamil intake, and IoT sensor feeds for automatic re-planning."),

    ("h2", "Segment 6 — Wrap Up (9:30–10:00)"),
    ("say", "ResQGrid AI turns fragmented reports into decisions: detect, understand, prioritize, allocate, approve, respond, re-plan. It is live, deployed, and installable on Android today. Because in an emergency, the AI should do the reading — and the human should make the call. Thank you. We are ready for your questions."),

    ("h1", "3. Live Demo Run-Sheet (5:00 inside the pitch)"),
    ("h3", "Pre-flight — before you are called on stage"),
    ("b", [
        "Laptop: Chrome open on the app login page, plus tabs for the Swagger API docs and the GitHub repo.",
        "Phone: the Android APK (or PWA) installed and logged in as the responder.",
        "Wake the backend at T-15 minutes: open https://resqgrid-api-9wns.onrender.com/docs — free-tier Render sleeps, and the first request can take about a minute.",
        "Demo logins printed on the cheat sheet (Section 8).",
        "Backup ready: a 2-minute screen recording of this exact flow, cued on the same laptop; phone hotspot as plan B for Wi-Fi; plan C is the local Docker Compose stack.",
    ]),
    ("h3", "The run-sheet"),
    ("table", {
        "headers": ["Step", "Time", "Do this", "Say this"],
        "rows": [
            ["1", "0:00–0:40", "Log in as dispatcher (dispatcher@resqgrid.local / dispatch123). Dashboard: KPI cards and the live map with 20 seeded incidents colour-coded by severity.", "“This is the operational picture one duty officer sees: every incident, resource and hazard on one map.”"],
            ["2", "0:40–1:40", "Create a new incident and narrate the citizen report: “Flooding on Kandy Road near the bridge, water rising fast, 12 people including an elderly person, medical need.” Run AI Triage. Point at the structured result: type, severity, people at risk, reason codes, confidence.", "“A panicked free-text report becomes structured triage in about two seconds — and it tells us why: rapid water rise, vulnerable person, medical need. Confidence is reported separately from severity.”"],
            ["3", "1:40–2:20", "Click Calculate Priority. Show the 0–100 score and the component breakdown: Life 30%, Medical 20%, People 15%, Environment 15%, Time 10%, Evidence 10%.", "“No black box: every point of this score decomposes into weighted, explainable factors. An operator can challenge any component.”"],
            ["4", "2:20–3:10", "Generate recommendations. Show the best-matched resource — type, capacity, ETA, compatibility reasons — plus the alternatives. Click Approve.", "“AI recommends — but nothing moves until a human approves. The approval, with the AI’s reasons, is written to an audit log.”"],
            ["5", "3:10–3:50", "Switch to the phone: the responder sees the assignment — Accept, then En Route. Back on desktop, mark a road blocked as a hazard. Show the system recalculating and proposing an alternative resource and route.", "“The field responder works from the same system on a phone. When the world changes — a road closes — the plan changes. Re-planning is automatic; re-approval is human.”"],
            ["6", "3:50–4:30", "Open the Command Assistant. Ask: “What are our top 3 priorities right now?” Then: “Which rescue units are free near Borella?”", "“The assistant answers from the live database — not from a script — and cites its sources. With an AI key it uses Qwen; without one it still answers from real data.”"],
            ["7", "4:30–5:00", "Hold up the phone: the same app, installed as a native Android APK via Capacitor.", "“Same system, packaged as a native Android app. A responder in the field needs nothing but a phone.”"],
        ],
    }),
    ("box", "Fallback rule: if the live backend stalls for more than about 20 seconds at any step, say “Production internet — let me switch to the recorded run of the same deployment” and play the backup video from that step. A smooth switch beats a frozen tab."),

    ("h1", "4. “How Did YOU Use AI to Build It?” — Your Explainable Answer"),
    ("p", "This is an AI buildathon: using AI heavily to build the product is expected. What earns points is explaining what you asked AI to do, how you guided and checked its output, and why you made your choices. If time is short, walk through this one concrete part."),
    ("h3", "The concrete part: the AI triage adapter"),
    ("b", [
        "What we asked: “Given this raw citizen report, extract incident type, severity, people at risk, vulnerable people, medical need, immediate needs, evidence quality and confidence — return strict JSON matching our schema; if information is missing, return null, never guess.”",
        "How we guided it: a strict system prompt carrying the JSON schema and few-shot examples; triage (Qwen) separated from assistant queries; low temperature for structured extraction.",
        "How we checked it: every response is validated with Pydantic v2 before it touches business logic; automated tests cover flood, fire and landslide reports; reason codes map to deterministic on-screen explanations; confidence is separated from severity so a low-confidence “critical” can never auto-dispatch.",
        "Why these choices: in an emergency product a wrong AI answer is worse than no answer — so all AI output is advisory, scored, logged and human-approved.",
        "Where AI helped us build: Qoder produced first-pass implementations of the adapter, the priority service and the assistant fallback. We owned the architecture — schemas, adapter pattern, human-in-the-loop rules — reviewed every generated file, and wrote the tests that pin the behaviour.",
    ]),
    ("box", "One-sentence version if a judge asks cold: “We used AI to build a schema-validated triage adapter — we specified the exact JSON contract and the null-not-guess rule, AI produced the implementation, and Pydantic validation plus tests guarantee the LLM can never write unvalidated data into the system.”"),

    ("h1", "5. Impact Numbers That Survive Follow-Up Questions"),
    ("table", {
        "headers": ["Claim", "Number", "What is behind it"],
        "rows": [
            ["Recent worst event", "330+ deaths", "BBC reporting on the November 2025 Sri Lanka floods and mudslides."],
            ["Urban-scale displacement", "300,000+ affected", "2016 monsoon floods — IOM and UN situation reports."],
            ["Recurrence", "Every year for 40 years", "World Bank GRADE rapid post-disaster damage assessment for Sri Lanka."],
            ["Triage speed", "Seconds vs minutes per report", "Live demo step 2: AI triage returns in about 2 seconds versus a phone-call triage of several minutes."],
            ["MVP running cost", "Near zero", "Deployed on free tiers: Vercel + Render + Supabase + Upstash."],
            ["Field hardware cost", "One Android phone", "The same web app installs as a PWA or a Capacitor-built APK."],
        ],
    }),
    ("p", "Feasibility path: pilot with one District DMC office during the next monsoon season; data stays in-country on PostgreSQL/PostGIS; the open-source stack avoids vendor lock-in; role-based access matches how an EOC already divides work."),

    ("h1", "6. Q&A — The 12 Questions You Will Probably Get"),
    ("nums", [
        ["Most teams built a chatbot. Why is this different?", "A chatbot answers one person. This is an operational system: many reporters in, one accountable decision-maker out — with prioritization, resource allocation and an audit trail. The assistant is one feature, not the product."],
        ["What happens when the AI is wrong?", "Three layers: the AI reports confidence and says “unknown” instead of guessing; the priority score is deterministic and challengeable factor by factor; and AI never dispatches — an authorized human approves every action, with everything logged."],
        ["Can it run without internet?", "The app is an installable PWA with a native Android APK, and locally the whole stack runs offline via Docker Compose. Roadmap: SMS and voice intake for zero-data scenarios."],
        ["Who exactly is the user?", "The duty officer at a Sri Lanka District Emergency Operations Centre during floods and landslides. Secondary users: dispatchers, field responders, and citizens submitting reports."],
        ["Why not just use WhatsApp groups?", "WhatsApp is where reports arrive — it does not structure, deduplicate, prioritize, match resources or audit decisions. ResQGrid can ingest from those channels; it replaces the spreadsheet-and-phone coordination on top of them."],
        ["Did you really build this in the buildathon period?", "Walk the Git history and the live deployment. We used AI coding tools heavily — that is expected here — but the architecture, schemas, safety rules and tests were our decisions, and we can explain any file."],
        ["Is the priority score scientifically validated?", "Honest answer: no. The weights are configurable starting values, clearly labelled as such. Calibrating them against historical DMC data is the first task of the pilot. That transparency is deliberate."],
        ["What if the LLM API is down mid-disaster?", "The provider adapter falls back: Qwen to Gemini to rule-based triage. The system degrades to a structured manual workflow instead of stopping. We would rather answer slower than answer wrong."],
        ["How does it scale from 20 demo incidents to a real district?", "The API is stateless and scales horizontally; PostGIS handles geospatial queries; Redis caches hot reads; AI triage is per-incident and parallel. The data model already separates incidents, resources, hazards and assignments."],
        ["Is citizen and victim data safe?", "JWT authentication with four role-based access levels; the assistant cannot expose restricted personal data; every consequential read and write is audit-logged; secrets live in environment variables only."],
        ["What about Sinhala and Tamil speakers?", "Intake is English today; multilingual LLM adapters make Sinhala and Tamil a roadmap item, not a rebuild. We state it openly as a next step."],
        ["What is real today versus plan?", "Real today: the deployed multi-user system — AI triage, explainable scoring, matching with human approval, hazard re-planning, assistant answers from live data, PWA and APK. Plan: DMC pilot, multilingual intake, IoT sensor feeds, SMS/voice channel, weight calibration with real data."],
    ]),

    ("h1", "7. The Night-Before and Morning-Of Checklist"),
    ("h3", "The organizers’ list (from the pitch guidelines)"),
    ("b", [
        "[ ] Demo tested and working within the last hour — re-test on the venue Wi-Fi before you go up.",
        "[ ] Backup video ready and cued on the same machine.",
        "[ ] Full run-through timed at 10 minutes or under — do two complete rehearsals tonight.",
        "[ ] Every team member knows their part — one speaks, one drives the demo, one watches the clock.",
        "[ ] You can clearly explain one concrete AI-built part (Section 4).",
        "[ ] You know your target user and one number behind your impact claim: duty officer at a DMC EOC; 330+ deaths in the Nov 2025 floods.",
    ]),
    ("h3", "ResQGrid-specific additions"),
    ("b", [
        "[ ] Wake the Render backend at T-15 minutes (open the /docs URL); keep a ping tab alive.",
        "[ ] Phone charged, APK logged in as responder1@resqgrid.local / respond123.",
        "[ ] HDMI or USB-C adapter, laptop charger, phone hotspot ready.",
        "[ ] Login page, Swagger docs and GitHub repo open in tabs; leave the dispatcher session logged in for judges afterwards.",
        "[ ] Water. Breathe. The demo is half the score — protect its time.",
    ]),

    ("h1", "8. One-Page Cheat Sheet (print this)"),
    ("h3", "The six scores at a glance"),
    ("p", "Problem 15 · AI centrality 25 · Working prototype 25 · Explainable AI 10 · Impact 15 · Demo clarity 10. Half the score is the AI plus the working build — that is why 5 of your 10 minutes are the demo."),
    ("h3", "Say these three numbers"),
    ("b", [
        "330+ deaths — November 2025 Sri Lanka floods and mudslides (BBC).",
        "300,000+ affected — 2016 monsoon floods (IOM/UN).",
        "~2 seconds — AI triage per report, live in the demo.",
    ]),
    ("h3", "Demo in seven lines"),
    ("nums", [
        ["Login as dispatcher — live map of 20 incidents."],
        ["New flood report, 12 people at risk — Run AI Triage — structured output + reason codes."],
        ["Calculate Priority — explainable weighted score."],
        ["Generate recommendations — approve one (human-in-the-loop + audit log)."],
        ["Responder accepts on the phone — mark a road blocked — system re-plans."],
        ["Assistant: “Top 3 priorities right now?” — answers from the live database."],
        ["Hold up the phone — same app as a native Android APK."],
    ]),
    ("h3", "Demo logins (judges can use the dispatcher account after your pitch)"),
    ("table", {
        "headers": ["Role", "Email", "Password", "Can do"],
        "rows": [
            ["Admin", "admin@resqgrid.local", "admin123", "Everything, including operational oversight"],
            ["Dispatcher", "dispatcher@resqgrid.local", "dispatch123", "Triage, priority, recommendations, approvals"],
            ["Responder", "responder1@resqgrid.local", "respond123", "Accept and complete assignments, update status"],
            ["Citizen", "citizen@resqgrid.local", "citizen123", "Submit and track incident reports"],
        ],
    }),
    ("h3", "If a judge asks how YOU used AI"),
    ("p", "“We specified the exact JSON contract and the null-not-guess rule; AI produced the triage-adapter implementation; Pydantic validation plus tests guarantee the LLM can never write unvalidated data into the system.”"),
    ("h3", "Fallback line (memorize)"),
    ("p", "“Production internet — let me switch to the recorded run of the same deployment.”"),
]

# ----------------------------------------------------------------- render ---
ACCENT = "C2410C"   # emergency orange
DARK = "1F2937"     # slate
GRAY = "4B5563"


def _hex(h):
    from docx.shared import RGBColor
    return RGBColor(int(h[0:2], 16), int(h[2:4], 16), int(h[4:6], 16))


def _shade(cell, color):
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), color)
    tcPr.append(shd)


def build_docx():
    from docx import Document
    from docx.shared import Pt, Cm, Inches
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    doc = Document()
    doc.core_properties.title = "ResQGrid AI — Pitch Day Playbook"
    doc.core_properties.author = "ResQGrid AI Team"

    for section in doc.sections:
        section.top_margin = Cm(2.0)
        section.bottom_margin = Cm(2.0)
        section.left_margin = Cm(2.2)
        section.right_margin = Cm(2.2)

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)

    def para(text, size=11, bold=False, italic=False, color=DARK, space_after=6,
             space_before=0, indent=None, align=None):
        p = doc.add_paragraph()
        r = p.add_run(text)
        r.font.size = Pt(size)
        r.bold = bold
        r.italic = italic
        r.font.color.rgb = _hex(color)
        p.paragraph_format.space_after = Pt(space_after)
        p.paragraph_format.space_before = Pt(space_before)
        if indent is not None:
            p.paragraph_format.left_indent = Cm(indent)
        if align:
            p.alignment = align
        return p

    def shade_para(p, fill):
        from docx.oxml.ns import qn
        from docx.oxml import OxmlElement
        pPr = p._p.get_or_add_pPr()
        shd = OxmlElement("w:shd")
        shd.set(qn("w:val"), "clear")
        shd.set(qn("w:fill"), fill)
        pPr.append(shd)

    def add_table(headers, rows):
        t = doc.add_table(rows=1, cols=len(headers))
        t.style = "Table Grid"
        hdr = t.rows[0].cells
        for i, h in enumerate(headers):
            hdr[i].text = ""
            r = hdr[i].paragraphs[0].add_run(h)
            r.bold = True
            r.font.size = Pt(10)
            r.font.color.rgb = _hex("FFFFFF")
            _shade(hdr[i], ACCENT)
        for row in rows:
            cells = t.add_row().cells
            for i, val in enumerate(row):
                cells[i].text = ""
                r = cells[i].paragraphs[0].add_run(val)
                r.font.size = Pt(10)
                if i == 0:
                    r.bold = True
        if len(headers) >= 2 and headers[1] in ("Pts", "Time", "Password"):
            t.columns[0].width = Cm(3.2)
        doc.add_paragraph().paragraph_format.space_after = Pt(2)

    first = True
    for kind, payload in CONTENT:
        if kind == "title":
            para(payload, size=26, bold=True, color=ACCENT, space_after=2)
        elif kind == "subtitle":
            para(payload, size=15, bold=True, color=DARK, space_after=10)
        elif kind == "meta":
            for line in payload:
                para(line, size=10.5, color=GRAY, space_after=2)
            doc.add_paragraph().paragraph_format.space_after = Pt(4)
        elif kind == "h1":
            if not first:
                doc.add_page_break()
            first = False
            para(payload, size=17, bold=True, color=ACCENT, space_before=4, space_after=8)
        elif kind == "h2":
            para(payload, size=13.5, bold=True, color=DARK, space_before=10, space_after=6)
        elif kind == "h3":
            para(payload, size=11.5, bold=True, color=ACCENT, space_before=8, space_after=4)
        elif kind == "p":
            para(payload, size=11)
        elif kind == "say":
            para("“" + payload + "”" if not payload.startswith("“") else payload,
                 size=11, italic=True, color=GRAY, indent=0.8, space_after=8)
        elif kind == "dir":
            para("[Stage direction: " + payload + "]", size=10, italic=True,
                 color=ACCENT, space_after=8)
        elif kind == "b":
            for item in payload:
                p = para("•  " + item, size=10.5, space_after=4, indent=0.4)
        elif kind == "nums":
            for idx, pair in enumerate(payload, 1):
                if len(pair) == 2:
                    para(f"Q{idx}. {pair[0]}", size=11, bold=True, space_before=6, space_after=2)
                    para(pair[1], size=10.5, color=GRAY, indent=0.5, space_after=6)
                else:
                    para(f"{idx}. {pair[0]}", size=10.5, space_after=3)
        elif kind == "table":
            add_table(payload["headers"], payload["rows"])
        elif kind == "box":
            p = para(payload, size=10.5, bold=True, color=DARK, space_before=6, space_after=8)
            shade_para(p, "FDE8DD")
        elif kind == "pb":
            doc.add_page_break()

    doc.save(DOCX_PATH)


def build_pdf():
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.units import cm
    from reportlab.lib import colors
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.enums import TA_LEFT
    from reportlab.platypus import (BaseDocTemplate, Frame, PageTemplate, Paragraph,
                                    Spacer, Table, TableStyle, PageBreak)
    from reportlab.pdfgen import canvas as pdfcanvas

    orange = colors.HexColor("#" + ACCENT)
    dark = colors.HexColor("#" + DARK)
    gray = colors.HexColor("#" + GRAY)
    light = colors.HexColor("#FDE8DD")
    grid = colors.HexColor("#9CA3AF")

    S = {
        "title": ParagraphStyle("title", fontName="Helvetica-Bold", fontSize=23,
                                textColor=orange, spaceAfter=4),
        "subtitle": ParagraphStyle("subtitle", fontName="Helvetica-Bold", fontSize=14,
                                   textColor=dark, spaceAfter=12),
        "meta": ParagraphStyle("meta", fontName="Helvetica", fontSize=9.5,
                               textColor=gray, spaceAfter=2),
        "h1": ParagraphStyle("h1", fontName="Helvetica-Bold", fontSize=15.5,
                             textColor=orange, spaceBefore=6, spaceAfter=8),
        "h2": ParagraphStyle("h2", fontName="Helvetica-Bold", fontSize=12.5,
                             textColor=dark, spaceBefore=10, spaceAfter=5),
        "h3": ParagraphStyle("h3", fontName="Helvetica-Bold", fontSize=11,
                             textColor=orange, spaceBefore=8, spaceAfter=4),
        "p": ParagraphStyle("p", fontName="Helvetica", fontSize=10.5, leading=14.5,
                            textColor=dark, spaceAfter=6),
        "say": ParagraphStyle("say", fontName="Helvetica-Oblique", fontSize=10.5,
                              leading=14.5, textColor=gray, leftIndent=16, spaceAfter=7),
        "dir": ParagraphStyle("dir", fontName="Helvetica-Oblique", fontSize=9.5,
                              textColor=orange, spaceAfter=7),
        "b": ParagraphStyle("b", fontName="Helvetica", fontSize=10.5, leading=14,
                            textColor=dark, leftIndent=12, spaceAfter=4),
        "num": ParagraphStyle("num", fontName="Helvetica", fontSize=10.5, leading=14,
                              textColor=dark, leftIndent=12, spaceAfter=3),
        "q": ParagraphStyle("q", fontName="Helvetica-Bold", fontSize=10.5, leading=14,
                            textColor=dark, spaceBefore=6, spaceAfter=2),
        "a": ParagraphStyle("a", fontName="Helvetica", fontSize=10.5, leading=14,
                            textColor=gray, leftIndent=14, spaceAfter=6),
        "box": ParagraphStyle("box", fontName="Helvetica-Bold", fontSize=10.5, leading=14,
                              textColor=dark, spaceBefore=6, spaceAfter=8),
        "cell": ParagraphStyle("cell", fontName="Helvetica", fontSize=9, leading=12,
                               textColor=dark),
        "cellb": ParagraphStyle("cellb", fontName="Helvetica-Bold", fontSize=9, leading=12,
                                textColor=dark),
        "cellh": ParagraphStyle("cellh", fontName="Helvetica-Bold", fontSize=9, leading=12,
                                textColor=colors.white),
    }

    def esc(t):
        return t.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")

    story = []
    pending_h1 = None

    def make_table(headers, rows, avail):
        n = len(headers)
        if n == 4 and headers[1] == "Pts":
            widths = [0.19 * avail, 0.05 * avail, 0.30 * avail, 0.46 * avail]
        elif n == 4 and headers[1] == "Time":
            widths = [0.05 * avail, 0.10 * avail, 0.50 * avail, 0.35 * avail]
        elif n == 4 and headers[1] == "Email":
            widths = [0.13 * avail, 0.32 * avail, 0.18 * avail, 0.37 * avail]
        elif n == 3:
            widths = [0.22 * avail, 0.18 * avail, 0.60 * avail]
        else:
            widths = [avail / n] * n
        data = [[Paragraph(esc(h), S["cellh"]) for h in headers]]
        for row in rows:
            data.append([Paragraph(esc(c), S["cellb" if i == 0 else "cell"])
                         for i, c in enumerate(row)])
        t = Table(data, colWidths=widths, repeatRows=1)
        t.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), orange),
            ("GRID", (0, 0), (-1, -1), 0.5, grid),
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F3F4F6")]),
            ("LEFTPADDING", (0, 0), (-1, -1), 5),
            ("RIGHTPADDING", (0, 0), (-1, -1), 5),
            ("TOPPADDING", (0, 0), (-1, -1), 4),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
        ]))
        return t

    avail = A4[0] - 4.4 * cm

    for kind, payload in CONTENT:
        if kind == "title":
            story.append(Paragraph(esc(payload), S["title"]))
        elif kind == "subtitle":
            story.append(Paragraph(esc(payload), S["subtitle"]))
        elif kind == "meta":
            for line in payload:
                story.append(Paragraph(esc(line), S["meta"]))
            story.append(Spacer(1, 8))
        elif kind == "h1":
            pending_h1 = payload
        elif kind == "h2":
            if pending_h1:
                story.append(PageBreak())
                story.append(Paragraph(esc(pending_h1), S["h1"]))
                pending_h1 = None
            story.append(Paragraph(esc(payload), S["h2"]))
        elif kind == "h3":
            if pending_h1:
                story.append(PageBreak())
                story.append(Paragraph(esc(pending_h1), S["h1"]))
                pending_h1 = None
            story.append(Paragraph(esc(payload), S["h3"]))
        elif kind == "p":
            if pending_h1:
                story.append(PageBreak())
                story.append(Paragraph(esc(pending_h1), S["h1"]))
                pending_h1 = None
            story.append(Paragraph(esc(payload), S["p"]))
        elif kind == "say":
            txt = payload if payload.startswith("“") else "“" + payload + "”"
            story.append(Paragraph(esc(txt), S["say"]))
        elif kind == "dir":
            story.append(Paragraph(esc("[Stage direction: " + payload + "]"), S["dir"]))
        elif kind == "b":
            if pending_h1:
                story.append(PageBreak())
                story.append(Paragraph(esc(pending_h1), S["h1"]))
                pending_h1 = None
            for item in payload:
                story.append(Paragraph("•  " + esc(item), S["b"]))
        elif kind == "nums":
            if pending_h1:
                story.append(PageBreak())
                story.append(Paragraph(esc(pending_h1), S["h1"]))
                pending_h1 = None
            for idx, pair in enumerate(payload, 1):
                if len(pair) == 2:
                    story.append(Paragraph(esc(f"Q{idx}. {pair[0]}"), S["q"]))
                    story.append(Paragraph(esc(pair[1]), S["a"]))
                else:
                    story.append(Paragraph(esc(f"{idx}. {pair[0]}"), S["num"]))
        elif kind == "table":
            if pending_h1:
                story.append(PageBreak())
                story.append(Paragraph(esc(pending_h1), S["h1"]))
                pending_h1 = None
            story.append(make_table(payload["headers"], payload["rows"], avail))
            story.append(Spacer(1, 8))
        elif kind == "box":
            t = Table([[Paragraph(esc(payload), S["box"])]], colWidths=[avail])
            t.setStyle(TableStyle([
                ("BACKGROUND", (0, 0), (-1, -1), light),
                ("BOX", (0, 0), (-1, -1), 0.75, orange),
                ("LEFTPADDING", (0, 0), (-1, -1), 8),
                ("RIGHTPADDING", (0, 0), (-1, -1), 8),
                ("TOPPADDING", (0, 0), (-1, -1), 6),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
            ]))
            story.append(t)

    if pending_h1:
        story.append(PageBreak())
        story.append(Paragraph(esc(pending_h1), S["h1"]))

    def footer(canv, doc_):
        canv.saveState()
        canv.setFont("Helvetica", 8)
        canv.setFillColor(gray)
        canv.drawString(2.2 * cm, 1.1 * cm, "ResQGrid AI — Pitch Day Playbook · 16 September 2026")
        canv.drawRightString(A4[0] - 2.2 * cm, 1.1 * cm, f"Page {canv.getPageNumber()}")
        canv.restoreState()

    pdf = BaseDocTemplate(PDF_PATH, pagesize=A4,
                          leftMargin=2.2 * cm, rightMargin=2.2 * cm,
                          topMargin=2.0 * cm, bottomMargin=2.0 * cm,
                          title="ResQGrid AI — Pitch Day Playbook",
                          author="ResQGrid AI Team")
    frame = Frame(pdf.leftMargin, pdf.bottomMargin, pdf.width, pdf.height, id="f")
    pdf.addPageTemplates([PageTemplate(id="p", frames=[frame], onPage=footer)])
    pdf.build(story)


if __name__ == "__main__":
    build_docx()
    build_pdf()
    for path in (DOCX_PATH, PDF_PATH):
        print(f"WROTE {path} ({os.path.getsize(path)} bytes)")
